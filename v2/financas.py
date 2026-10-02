"""Family 2 (Planner Financeiro 2027): printable planner, 52-week challenge (bump), debt guide + spreadsheet (upsell)."""
import calendar
import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from design import (FINANCAS, Flow, H, M, W, C, coin, cover, draw_para, header_footer, para_height, rich,
                    rounded_box, style, writing_lines, mm)

T = FINANCAS
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro",
         "Novembro", "Dezembro"]
DIAS = ["S", "T", "Q", "Q", "S", "S", "D"]


def easter(year):
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l_ = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l_) // 451
    month = (h + l_ - 7 * m + 114) // 31
    day = (h + l_ - 7 * m + 114) % 31 + 1
    return dt.date(year, month, day)


def feriados(year=2027):
    e = easter(year)
    return [
        (dt.date(year, 1, 1), "Confraternização Universal", True),
        (e - dt.timedelta(days=48), "Carnaval (ponto facultativo)", False),
        (e - dt.timedelta(days=47), "Carnaval (ponto facultativo)", False),
        (e - dt.timedelta(days=2), "Sexta-feira Santa", True),
        (dt.date(year, 4, 21), "Tiradentes", True),
        (dt.date(year, 5, 1), "Dia do Trabalho", True),
        (e + dt.timedelta(days=60), "Corpus Christi (ponto facultativo)", False),
        (dt.date(year, 9, 7), "Independência do Brasil", True),
        (dt.date(year, 10, 12), "Nossa Senhora Aparecida", True),
        (dt.date(year, 11, 2), "Finados", True),
        (dt.date(year, 11, 15), "Proclamação da República", True),
        (dt.date(year, 11, 20), "Dia Nacional de Zumbi e da Consciência Negra", True),
        (dt.date(year, 12, 25), "Natal", True),
    ]


def page_bg(c, page_no, title="Planner Financeiro 2027"):
    c.setFillColor(T.paper)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    header_footer(c, T, title, page_no)


def kicker_title(c, kicker, title, y=None):
    y = y or H - 22 * mm
    c.setFillColor(T.accent)
    c.setFont("Body-Bold", 9)
    c.drawString(M, y, kicker.upper())
    return draw_para(c, title, style("kt", fontName="Serif-Bold", fontSize=22, leading=26, textColor=T.primary),
                     M, y - 3 * mm, W - 2 * M) - 4 * mm


def table(c, x, y, widths, headers, rows, row_h=6.2 * mm, head_fill=None, font_size=9.2, zebra=True):
    """Draw a simple fill-in table. rows: list of lists of strings. Returns y after the table."""
    total_w = sum(widths)
    c.setFillColor(head_fill or T.primary)
    c.rect(x, y - row_h, total_w, row_h, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont("Body-Bold", font_size)
    cx = x
    for wdt, h_ in zip(widths, headers):
        c.drawString(cx + 2 * mm, y - row_h + 2 * mm, h_)
        cx += wdt
    y -= row_h
    c.setFont("Body", font_size)
    for i, r in enumerate(rows):
        if zebra and i % 2 == 0:
            c.setFillColor(T.soft)
            c.rect(x, y - row_h, total_w, row_h, stroke=0, fill=1)
        c.setStrokeColor(C("#D5DDD7"))
        c.setLineWidth(0.4)
        c.line(x, y - row_h, x + total_w, y - row_h)
        cx = x
        c.setFillColor(T.ink)
        for wdt, val in zip(widths, r):
            if val == "[]":
                c.setStrokeColor(T.primary)
                c.rect(cx + wdt / 2 - 1.8 * mm, y - row_h + 1.4 * mm, 3.6 * mm, 3.6 * mm, stroke=1, fill=0)
            else:
                c.drawString(cx + 2 * mm, y - row_h + 2 * mm, str(val))
            cx += wdt
        y -= row_h
    c.setStrokeColor(C("#B9C7BD"))
    cx = x
    for wdt in widths[:-1]:
        cx += wdt
        c.line(cx, y, cx, y + row_h * (len(rows) + 1))
    return y


def mini_month(c, x, y, year, month, hol_days, w=56 * mm):
    cell = w / 7
    c.setFillColor(T.primary)
    c.setFont("Body-Bold", 10)
    c.drawString(x, y, MESES[month - 1])
    y -= 5.5 * mm
    c.setFont("Body-Bold", 7.5)
    c.setFillColor(T.accent)
    for i, d in enumerate(DIAS):
        c.drawCentredString(x + cell * i + cell / 2, y, d)
    y -= 5 * mm
    c.setFont("Body", 8.5)
    for week in calendar.Calendar(firstweekday=0).monthdayscalendar(year, month):
        for i, d in enumerate(week):
            if not d:
                continue
            cx = x + cell * i + cell / 2
            if (month, d) in hol_days:
                c.setFillColor(T.secondary)
                c.circle(cx, y + 1.1 * mm, 2.6 * mm, stroke=0, fill=1)
                c.setFillColor(T.primary)
                c.setFont("Body-Bold", 8.5)
            else:
                c.setFillColor(T.ink if i < 5 else T.accent)
                c.setFont("Body", 8.5)
            c.drawCentredString(cx, y, str(d))
        y -= 5 * mm


def build_planner(data, out):
    c = canvas.Canvas(out, pagesize=(W, H))
    c.setTitle("Planner Financeiro 2027")
    cover(c, T, "Planner Financeiro 2027", "Orçamento, contas, dívidas e metas: o ano inteiro organizado no papel",
          "Para imprimir · 12 meses · calendário 2027 com feriados")
    c.showPage()
    f = Flow(c, T, "Planner Financeiro 2027", start_page=2)
    f.new_page()
    f.title_block("Comece por aqui", data.get("intro_title", "Bem-vindo ao seu 2027"))
    for p in data.get("intro_paragraphs", []):
        f.para(p)
    f.heading("Como usar o planner")
    f.bullets(data.get("how_to_use", []), numbered=True)
    f.para("Conteúdo educativo para organização pessoal. Não é recomendação de investimento.", f.small)
    c.showPage()
    page = f.page + 1
    hol = feriados(2027)
    hol_days = {(d.month, d.day) for d, _, _ in hol}
    for half in (0, 1):
        page_bg(c, page)
        y = kicker_title(c, "Calendário 2027", "Primeiro semestre" if half == 0 else "Segundo semestre")
        for k in range(6):
            mth = half * 6 + k + 1
            col, row = k % 3, k // 3
            mini_month(c, M + col * 60 * mm, y - row * 52 * mm, 2027, mth, hol_days)
        yy = y - 2 * 52 * mm - 4 * mm
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 10)
        c.drawString(M, yy, "Feriados nacionais e datas importantes")
        yy -= 6 * mm
        c.setFont("Body", 9.5)
        for d, name, _ in [h_ for h_ in hol if (h_[0].month <= 6) == (half == 0)]:
            c.setFillColor(T.ink)
            c.drawString(M, yy, f"{d.day:02d}/{d.month:02d}  ·  {name}")
            yy -= 5.2 * mm
        c.setFillColor(C("#7A7A7A"))
        c.setFont("Body", 8)
        c.drawString(M, 16 * mm, "Feriados estaduais e municipais variam: anote os da sua cidade.")
        c.showPage()
        page += 1
    si = data.get("section_intros", {})
    # goals
    page_bg(c, page)
    y = kicker_title(c, "Metas", "Minhas metas para 2027")
    y = draw_para(c, si.get("goals", ""), style("si", fontSize=11, leading=16, textColor=T.ink), M, y, W - 2 * M) - 5 * mm
    for g in range(1, 4):
        rounded_box(c, M, y - 62 * mm, W - 2 * M, 60 * mm, T.soft, r=6)
        c.setFillColor(T.primary)
        c.setFont("Serif-Bold", 14)
        c.drawString(M + 6 * mm, y - 9 * mm, f"Meta {g}")
        c.setFont("Body", 9.5)
        c.setFillColor(T.ink)
        for j, lab in enumerate(["O quê:", "Quanto (R$):", "Até quando:", "Por que é importante:", "Primeiro passo:"]):
            yy = y - 18 * mm - j * 9 * mm
            c.drawString(M + 6 * mm, yy, lab)
            c.setStrokeColor(C("#B9C7BD"))
            c.line(M + 42 * mm, yy - 1 * mm, W - M - 6 * mm, yy - 1 * mm)
        y -= 66 * mm
    c.showPage()
    page += 1
    # annual overview
    page_bg(c, page)
    y = kicker_title(c, "Visão do ano", "Meu ano em números")
    y = table(c, M, y, [32 * mm, 31 * mm, 31 * mm, 31 * mm, 49 * mm],
              ["Mês", "Receitas (R$)", "Despesas (R$)", "Guardado (R$)", "Observações"],
              [[m_, "", "", "", ""] for m_ in MESES] + [["Total", "", "", "", ""]], row_h=10 * mm, font_size=10)
    c.showPage()
    page += 1
    months = data.get("months", [])
    inc, fix, var = data.get("categories_income", []), data.get("categories_fixed", []), data.get("categories_variable", [])
    for mi, mname in enumerate(MESES):
        md = months[mi] if mi < len(months) else {"theme": "", "tip": "", "mini_challenge": ""}
        # budget page
        page_bg(c, page)
        c.setFillColor(T.primary)
        c.rect(0, H - 38 * mm, W, 30 * mm, stroke=0, fill=1)
        c.setFillColor(T.accent)
        c.setFont("Body-Bold", 9)
        c.drawString(M, H - 16 * mm, f"MÊS {mi + 1:02d} · ORÇAMENTO")
        c.setFillColor(colors.white)
        c.setFont("Serif-Black", 26)
        c.drawString(M, H - 27 * mm, f"{mname} 2027")
        c.setFont("Serif-It", 11)
        c.drawString(M, H - 34 * mm, md.get("theme", "")[:80])
        coin(c, W - M - 10 * mm, H - 23 * mm, 9 * mm, T.accent, T.primary)
        y = H - 44 * mm
        cols = [62 * mm, 34 * mm, 34 * mm, 44 * mm]
        y = table(c, M, y, cols, ["Receitas", "Previsto", "Realizado", "Observação"], [[x_, "", "", ""] for x_ in inc[:8]],
                  row_h=5.6 * mm, font_size=8.8) - 3 * mm
        y = table(c, M, y, [62 * mm, 30 * mm, 30 * mm, 24 * mm, 28 * mm], ["Despesas fixas", "Previsto", "Realizado",
                                                                         "Vence", "Pago"],
                  [[x_, "", "", "", "[]"] for x_ in fix[:14]], row_h=5.6 * mm, font_size=8.8) - 3 * mm
        y = table(c, M, y, cols, ["Despesas variáveis", "Previsto", "Realizado", "Observação"],
                  [[x_, "", "", ""] for x_ in var[:14]], row_h=5.6 * mm, font_size=8.8) - 4 * mm
        rounded_box(c, M, max(y - 16 * mm, 14 * mm), W - 2 * M, 14 * mm, T.soft, r=5)
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 10)
        c.drawString(M + 5 * mm, max(y - 16 * mm, 14 * mm) + 5 * mm,
                     "Receitas  R$ ________   –   Despesas  R$ ________   =   Saldo do mês  R$ ________")
        c.showPage()
        page += 1
        # daily spending page
        page_bg(c, page)
        y = kicker_title(c, f"{mname} · controle diário", "Para onde foi o meu dinheiro?")
        ndays = calendar.monthrange(2027, mi + 1)[1]
        y = table(c, M, y, [16 * mm, 86 * mm, 40 * mm, 32 * mm], ["Dia", "Descrição", "Categoria", "Valor (R$)"],
                  [[f"{d:02d}", "", "", ""] for d in range(1, ndays + 1)], row_h=7.2 * mm, font_size=8.5)
        c.showPage()
        page += 1
        # review page
        page_bg(c, page)
        y = kicker_title(c, f"{mname} · revisão do mês", "Como foi o mês?")
        y = draw_para(c, si.get("review", ""), style("rv", fontSize=10.5, leading=15, textColor=T.ink), M, y,
                      W - 2 * M) - 4 * mm
        for q in ["O que deu certo com o dinheiro este mês?", "Onde gastei mais do que planejei, e por quê?",
                  "O que vou fazer diferente no próximo mês?"]:
            c.setFillColor(T.primary)
            c.setFont("Body-Bold", 10.5)
            c.drawString(M, y, q)
            y = writing_lines(c, M, y - 1 * mm, W - 2 * M, 3, 8 * mm) - 6 * mm
        tip = md.get("tip", "")
        st = style("tip", fontSize=10.5, leading=15, textColor=T.ink)
        hb = para_height(tip, st, W - 2 * M - 14 * mm) + 14 * mm
        rounded_box(c, M, y - hb, W - 2 * M, hb, T.soft, r=6)
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 8.5)
        c.drawString(M + 7 * mm, y - 7 * mm, f"DICA DE {mname.upper()}")
        draw_para(c, tip, st, M + 7 * mm, y - 9.5 * mm, W - 2 * M - 14 * mm)
        y -= hb + 6 * mm
        rounded_box(c, M, y - 18 * mm, W - 2 * M, 18 * mm, C("#FFF6DD"), T.accent, r=6)
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 8.5)
        c.drawString(M + 7 * mm, y - 7 * mm, "MINI DESAFIO DO MÊS")
        c.setStrokeColor(T.primary)
        c.rect(W - M - 12 * mm, y - 12 * mm, 5 * mm, 5 * mm, stroke=1, fill=0)
        draw_para(c, md.get("mini_challenge", ""), style("mc", fontSize=10.5, leading=14, textColor=T.ink),
                  M + 7 * mm, y - 9 * mm, W - 2 * M - 26 * mm)
        c.showPage()
        page += 1
    # trackers
    page_bg(c, page)
    y = kicker_title(c, "Contas do ano", "Contas a pagar: marque quando pagar")
    y = draw_para(c, si.get("bills", ""), style("b", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 3 * mm
    mw = (W - 2 * M - 44 * mm) / 12
    table(c, M, y, [44 * mm] + [mw] * 12, ["Conta"] + [m_[:3] for m_ in MESES],
          [[x_] + ["[]"] * 12 for x_ in fix[:14]] + [[""] + ["[]"] * 12 for _ in range(6)], row_h=9 * mm, font_size=8.5)
    c.showPage()
    page += 1
    page_bg(c, page)
    y = kicker_title(c, "Dívidas", "Meu plano para sair das dívidas")
    y = draw_para(c, si.get("debts", ""), style("d", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 3 * mm
    y = table(c, M, y, [46 * mm, 30 * mm, 26 * mm, 28 * mm, 22 * mm, 22 * mm],
              ["Credor", "Saldo (R$)", "Juros/mês", "Parcela (R$)", "Ordem", "Quitada"],
              [["", "", "", "", "", "[]"] for _ in range(8)], row_h=9 * mm) - 8 * mm
    c.setFillColor(T.primary)
    c.setFont("Body-Bold", 11)
    c.drawString(M, y, "Progresso: pinte um quadrado a cada 10% quitado")
    y -= 8 * mm
    for k in range(5):
        c.setFillColor(T.ink)
        c.setFont("Body", 9.5)
        c.drawString(M, y - 4 * mm, "Dívida: ______________")
        for j in range(10):
            c.setStrokeColor(T.primary)
            c.rect(M + 48 * mm + j * 12.5 * mm, y - 6 * mm, 10 * mm, 7 * mm, stroke=1, fill=0)
        y -= 12 * mm
    c.showPage()
    page += 1
    page_bg(c, page)
    y = kicker_title(c, "Segurança", "Reserva de emergência")
    y = draw_para(c, si.get("emergency", ""), style("e", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 4 * mm
    c.setFont("Body", 10.5)
    c.setFillColor(T.ink)
    c.drawString(M, y, "Minha meta: R$ ____________    (sugestão: comece pequeno e vá aumentando)")
    y -= 10 * mm
    for j in range(20):
        col, row = j % 5, j // 5
        x = M + col * 35 * mm
        yy = y - row * 26 * mm
        rounded_box(c, x, yy - 22 * mm, 30 * mm, 22 * mm, colors.white, T.primary, r=5)
        c.setFillColor(T.accent)
        c.setFont("Body-Bold", 11)
        c.drawCentredString(x + 15 * mm, yy - 9 * mm, f"{(j + 1) * 5}%")
        c.setFillColor(C("#8A8A8A"))
        c.setFont("Body", 8)
        c.drawCentredString(x + 15 * mm, yy - 16 * mm, "R$ ________")
    c.showPage()
    page += 1
    page_bg(c, page)
    y = kicker_title(c, "Sonhos com data", "Metas de poupança")
    y = draw_para(c, si.get("savings", ""), style("s", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 4 * mm
    for g in range(3):
        c.setFillColor(T.primary)
        c.setFont("Serif-Bold", 13)
        c.drawString(M, y, f"Objetivo {g + 1}: ______________________   Valor: R$ ________   Até: ____/____")
        y -= 7 * mm
        for j in range(20):
            coin(c, M + 5 * mm + (j % 10) * 17 * mm, y - 5 * mm - (j // 10) * 14 * mm, 5.5 * mm, colors.white, T.primary)
        y -= 34 * mm
    c.showPage()
    page += 1
    page_bg(c, page)
    y = kicker_title(c, "Planejar antes de comprar", "Compras planejadas")
    y = draw_para(c, si.get("purchases", ""), style("p", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 3 * mm
    table(c, M, y, [60 * mm, 30 * mm, 30 * mm, 30 * mm, 24 * mm], ["O quê", "Preço (R$)", "Guardado", "Quando", "Comprei"],
          [["", "", "", "", "[]"] for _ in range(18)], row_h=10 * mm)
    c.showPage()
    page += 1
    page_bg(c, page)
    y = kicker_title(c, "Fechamento", "Resumo do meu 2027")
    y = draw_para(c, si.get("annual", ""), style("a", fontSize=10.5, leading=15, textColor=T.ink), M, y, W - 2 * M) - 4 * mm
    for q in ["Quanto eu ganhei no ano?", "Quanto eu guardei?", "Quais dívidas eu quitei?", "Do que eu mais me orgulho?",
              "Minha principal meta para 2028:"]:
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 11)
        c.drawString(M, y, q)
        y = writing_lines(c, M, y - 1 * mm, W - 2 * M, 2, 8.5 * mm) - 6 * mm
    c.showPage()
    page += 1
    page_bg(c, page)
    kicker_title(c, "Anotações", "Notas")
    writing_lines(c, M, H - 40 * mm, W - 2 * M, 26, 9 * mm)
    c.showPage()
    c.save()
    return out


def build_challenge(data, out):
    c = canvas.Canvas(out, pagesize=(W, H))
    c.setTitle("Desafio das 52 Semanas: Kit Poupança")
    cover(c, T, "Desafio das 52 Semanas", "Kit Poupança: três versões do desafio e 52 mini desafios para o seu 2027",
          "Bônus do Planner Financeiro 2027 · para imprimir")
    c.showPage()
    f = Flow(c, T, "Desafio das 52 Semanas", start_page=2)
    f.new_page()
    f.title_block("Como funciona", "Guarde um pouco toda semana")
    for p in data.get("challenge_intro_paragraphs", []):
        f.para(p)
    f.heading("Regras do desafio")
    f.bullets(data.get("challenge_rules", []), numbered=True)
    f.heading("Escolha a sua versão")
    for v in data.get("challenge_versions", []):
        f.para(f"<b>{rich(v.get('name', ''))}</b>: {rich(v.get('description', ''))}", raw=True)
    c.showPage()
    page = f.page + 1
    versions = [("Clássico", [i for i in range(1, 53)]), ("Turbo", [5 * i for i in range(1, 53)]),
                ("Invertido", [53 - i for i in range(1, 53)])]
    for name, amounts in versions:
        page_bg(c, page, "Desafio das 52 Semanas")
        total = sum(amounts)
        y = kicker_title(c, "Desafio das 52 semanas", f"Versão {name}")
        c.setFillColor(T.ink)
        c.setFont("Body", 10.5)
        c.drawString(M, y, f"Total ao final das 52 semanas: R$ {total:,}".replace(",", ".") + ",00")
        y -= 8 * mm
        cw_, ch_ = (W - 2 * M - 9 * mm) / 4, 15.2 * mm
        for i, a in enumerate(amounts):
            col, row = i % 4, i // 4
            x = M + col * (cw_ + 3 * mm)
            yy = y - row * (ch_ + 1.8 * mm)
            rounded_box(c, x, yy - ch_, cw_, ch_, colors.white, T.primary, r=4, lw=0.8)
            c.setFillColor(T.accent)
            c.setFont("Body-Bold", 8)
            c.drawString(x + 3 * mm, yy - 5 * mm, f"SEMANA {i + 1:02d}")
            c.setFillColor(T.primary)
            c.setFont("Serif-Bold", 13)
            c.drawString(x + 3 * mm, yy - 11.5 * mm, f"R$ {a},00")
            c.setStrokeColor(T.primary)
            c.rect(x + cw_ - 8 * mm, yy - 11 * mm, 5 * mm, 5 * mm, stroke=1, fill=0)
        c.showPage()
        page += 1
    page_bg(c, page, "Desafio das 52 Semanas")
    y = kicker_title(c, "Um por semana", "52 mini desafios para gastar menos")
    minis = data.get("mini_challenges", [])[:52]
    st = style("mini", fontSize=8.6, leading=10.6, textColor=T.ink)
    colw = (W - 2 * M - 8 * mm) / 2
    for i, mtxt in enumerate(minis):
        col, row = i // 26, i % 26
        x = M + col * (colw + 8 * mm)
        yy = y - row * 9.1 * mm
        c.setStrokeColor(T.primary)
        c.rect(x, yy - 3.8 * mm, 3.6 * mm, 3.6 * mm, stroke=1, fill=0)
        draw_para(c, f"{i + 1}. {mtxt}", st, x + 6 * mm, yy, colw - 6 * mm)
    c.showPage()
    page += 1
    page_bg(c, page, "Desafio das 52 Semanas")
    y = kicker_title(c, "Motivação", "Para que eu estou guardando?")
    for q in ["Meu objetivo com este dinheiro:", "Por que isso é importante para mim:", "Como vou comemorar quando terminar:"]:
        c.setFillColor(T.primary)
        c.setFont("Body-Bold", 11)
        c.drawString(M, y, q)
        y = writing_lines(c, M, y - 1 * mm, W - 2 * M, 4, 9 * mm) - 8 * mm
    c.showPage()
    c.save()
    return out


def build_guide(data, out):
    c = canvas.Canvas(out, pagesize=(W, H))
    c.setTitle(data.get("title", "Guia Saia das Dívidas em 2027"))
    cover(c, T, data.get("title", "Saia das Dívidas em 2027"), data.get("subtitle", ""),
          "Guia prático · bônus da Planilha Financeira 2027")
    c.showPage()
    f = Flow(c, T, data.get("title", "Saia das Dívidas"), start_page=2)
    f.new_page()
    f.title_block("Introdução", "Antes de tudo: respira")
    for p in data.get("intro_paragraphs", []):
        f.para(p)
    for i, ch in enumerate(data.get("chapters", []), 1):
        f.new_page()
        f.title_block(f"Capítulo {i}", ch.get("title", ""))
        for p in ch.get("paragraphs", []):
            f.para(p)
        if ch.get("steps"):
            f.heading("Passo a passo")
            f.bullets(ch["steps"], numbered=True)
        for co in ch.get("callouts", []):
            kind = co.get("kind", "tip")
            lab, fill, edge = {"tip": ("Dica", T.soft, T.primary), "warning": ("Atenção", C("#FBE7E5"), C("#A4161A")),
                               "example": ("Exemplo", C("#FFF6DD"), T.accent)}.get(kind, ("Dica", T.soft, T.primary))
            f.boxed(lab, co.get("text", ""), fill, edge)
        for sc in ch.get("scripts", []):
            f.boxed("Modelo de mensagem · " + sc.get("title", ""), sc.get("text", ""), C("#EAF4EE"), T.primary)
    f.new_page()
    f.title_block("Para fechar", "Um passo de cada vez")
    for p in data.get("closing_paragraphs", []):
        f.para(p)
    f.boxed("Aviso", data.get("disclaimer", ""), C("#F2F2F2"), C("#6B6B6B"))
    c.showPage()
    c.save()
    return out


# ---------------------------------------------------------------- spreadsheet
FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF3B0")
HEAD = PatternFill("solid", fgColor="1F5135")
SOFT = PatternFill("solid", fgColor="E6F2EA")
THIN = Side(style="thin", color="B9C7BD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BRL = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
SHORT = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]


def _hdr(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = Alignment(vertical="center")


def build_spreadsheet(data, out):
    inc, fix, var = data.get("categories_income", []), data.get("categories_fixed", []), data.get("categories_variable", [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Como usar"
    ws["A1"] = "Planilha Financeira 2027"
    ws["A1"].font = Font(name=FONT, bold=True, size=18, color="1F5135")
    ws["A3"] = "Legenda: células AMARELAS são para você preencher. As demais calculam sozinhas."
    ws["A3"].font = Font(name=FONT, bold=True)
    ws["A4"].fill = YELLOW
    ws["B4"] = "← preencha aqui"
    for i, step in enumerate(data.get("spreadsheet_instructions", []), 6):
        ws.cell(row=i, column=1, value=f"{i - 5}. {step}").font = Font(name=FONT)
    r = 7 + len(data.get("spreadsheet_instructions", []))
    ws.cell(row=r, column=1, value="Conteúdo educativo para organização pessoal. Não é recomendação de investimento.").font = \
        Font(name=FONT, italic=True, color="6B6B6B")
    ws.column_dimensions["A"].width = 110

    cat = wb.create_sheet("Categorias")
    _hdr(cat, 1, ["Receitas", "Despesas fixas", "Despesas variáveis"])
    for i in range(14):
        for col, lst in enumerate([inc, fix, var], 1):
            cell = cat.cell(row=2 + i, column=col, value=lst[i] if i < len(lst) else None)
            cell.fill = YELLOW
            cell.font = Font(name=FONT)
            cell.border = BOX
    for col in "ABC":
        cat.column_dimensions[col].width = 30
    cat["E1"] = "Edite os nomes das categorias aqui: os meses atualizam sozinhos."
    cat["E1"].font = Font(name=FONT, italic=True)

    # month sheets
    blocks = [("Receitas", "A", 8), ("Despesas fixas", "B", 14), ("Despesas variáveis", "C", 14)]
    totals_rows = {}
    for mi, sh in enumerate(SHORT):
        m = wb.create_sheet(sh)
        m["A1"] = f"{MESES[mi]} 2027"
        m["A1"].font = Font(name=FONT, bold=True, size=16, color="1F5135")
        row = 3
        tots = {}
        for title, col, n in blocks:
            _hdr(m, row, [title, "Previsto (R$)", "Realizado (R$)", "Diferença (R$)"])
            first = row + 1
            for i in range(n):
                rr = row + 1 + i
                m.cell(row=rr, column=1, value=f'=IF(Categorias!{col}{2 + i}="","",Categorias!{col}{2 + i})').font = Font(name=FONT)
                for cc in (2, 3):
                    cell = m.cell(row=rr, column=cc)
                    cell.fill = YELLOW
                    cell.number_format = BRL
                    cell.font = Font(name=FONT, color="0000FF")
                d = m.cell(row=rr, column=4, value=f"=C{rr}-B{rr}")
                d.number_format = BRL
                d.font = Font(name=FONT)
                for cc in range(1, 5):
                    m.cell(row=rr, column=cc).border = BOX
            last = row + n
            tr = last + 1
            m.cell(row=tr, column=1, value=f"Total {title.lower()}").font = Font(name=FONT, bold=True)
            for cc, L in ((2, "B"), (3, "C"), (4, "D")):
                cell = m.cell(row=tr, column=cc, value=f"=SUM({L}{first}:{L}{last})")
                cell.number_format = BRL
                cell.font = Font(name=FONT, bold=True)
                cell.fill = SOFT
            tots[title] = tr
            row = tr + 2
        sr = row
        m.cell(row=sr, column=1, value="Saldo do mês (receitas - despesas)").font = Font(name=FONT, bold=True, size=12)
        for cc, L in ((2, "B"), (3, "C")):
            cell = m.cell(row=sr, column=cc,
                          value=f"={L}{tots['Receitas']}-{L}{tots['Despesas fixas']}-{L}{tots['Despesas variáveis']}")
            cell.number_format = BRL
            cell.font = Font(name=FONT, bold=True, size=12)
            cell.fill = SOFT
        m.cell(row=sr + 1, column=1, value="Quanto guardei este mês").font = Font(name=FONT)
        g = m.cell(row=sr + 1, column=3)
        g.fill = YELLOW
        g.number_format = BRL
        tots["saldo"] = sr
        tots["guardado"] = sr + 1
        m.column_dimensions["A"].width = 34
        for L in "BCD":
            m.column_dimensions[L].width = 17
        if mi == 0:
            m["B4"] = 3000
            m["C4"] = 3000
            m["B4"].comment = Comment("Exemplo: apague e coloque o seu valor.", "Planilha")
        totals_rows = tots

    res = wb.create_sheet("Resumo", 1)
    res["A1"] = "Resumo de 2027"
    res["A1"].font = Font(name=FONT, bold=True, size=16, color="1F5135")
    _hdr(res, 3, ["Mês", "Receitas (R$)", "Despesas (R$)", "Saldo (R$)", "Guardado (R$)"])
    t = totals_rows
    for i, sh in enumerate(SHORT):
        rr = 4 + i
        res.cell(row=rr, column=1, value=MESES[i]).font = Font(name=FONT)
        vals = [f"='{sh}'!C{t['Receitas']}", f"='{sh}'!C{t['Despesas fixas']}+'{sh}'!C{t['Despesas variáveis']}",
                f"='{sh}'!C{t['saldo']}", f"='{sh}'!C{t['guardado']}"]
        for cc, v in enumerate(vals, 2):
            cell = res.cell(row=rr, column=cc, value=v)
            cell.number_format = BRL
            cell.font = Font(name=FONT, color="008000")
            cell.border = BOX
    res.cell(row=16, column=1, value="Total do ano").font = Font(name=FONT, bold=True)
    for cc, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        cell = res.cell(row=16, column=cc, value=f"=SUM({L}4:{L}15)")
        cell.number_format = BRL
        cell.font = Font(name=FONT, bold=True)
        cell.fill = SOFT
    res.cell(row=17, column=1, value="Média por mês").font = Font(name=FONT)
    for cc, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        cell = res.cell(row=17, column=cc, value=f"=AVERAGE({L}4:{L}15)")
        cell.number_format = BRL
        cell.font = Font(name=FONT)
    for L, wdt in zip("ABCDE", (16, 18, 18, 18, 18)):
        res.column_dimensions[L].width = wdt
    ch = BarChart()
    ch.title = "Receitas x Despesas"
    ch.y_axis.title = "R$"
    ch.add_data(Reference(res, min_col=2, min_row=3, max_col=3, max_row=15), titles_from_data=True)
    ch.set_categories(Reference(res, min_col=1, min_row=4, max_row=15))
    ch.height, ch.width = 9, 18
    res.add_chart(ch, "G3")

    debt = wb.create_sheet("Dívidas")
    debt["A1"] = "Minhas dívidas: bola de neve x avalanche"
    debt["A1"].font = Font(name=FONT, bold=True, size=16, color="1F5135")
    debt["A2"] = ("Bola de neve: pague primeiro a MENOR dívida. Avalanche: pague primeiro a de MAIOR juros. "
                  "Juros ao mês em % (ex.: 0,08 = 8%).")
    debt["A2"].font = Font(name=FONT, italic=True)
    _hdr(debt, 4, ["Credor", "Saldo devedor (R$)", "Juros ao mês (%)", "Parcela que consigo pagar (R$)",
                   "Ordem bola de neve", "Ordem avalanche", "Meses para quitar", "Total de juros (R$)"])
    for i in range(10):
        rr = 5 + i
        for cc in (1, 2, 3, 4):
            cell = debt.cell(row=rr, column=cc)
            cell.fill = YELLOW
            cell.font = Font(name=FONT, color="0000FF")
            cell.border = BOX
        debt.cell(row=rr, column=2).number_format = BRL
        debt.cell(row=rr, column=3).number_format = "0.00%"
        debt.cell(row=rr, column=4).number_format = BRL
        debt.cell(row=rr, column=5, value=f'=IF(B{rr}="","",RANK(B{rr},$B$5:$B$14,1))')
        debt.cell(row=rr, column=6, value=f'=IF(C{rr}="","",RANK(C{rr},$C$5:$C$14,0))')
        debt.cell(row=rr, column=7, value=f'=IF(OR(B{rr}="",D{rr}=""),"",IFERROR(ROUNDUP(NPER(C{rr},-D{rr},B{rr}),0),"Parcela não cobre os juros"))')
        tj = debt.cell(row=rr, column=8, value=f'=IF(ISNUMBER(G{rr}),G{rr}*D{rr}-B{rr},"")')
        tj.number_format = BRL
        for cc in (5, 6, 7, 8):
            debt.cell(row=rr, column=cc).font = Font(name=FONT)
            debt.cell(row=rr, column=cc).border = BOX
    debt["A5"], debt["B5"], debt["C5"], debt["D5"] = "Exemplo: cartão da loja", 1200, 0.08, 150
    debt["A5"].comment = Comment("Linha de exemplo: substitua pelos seus dados.", "Planilha")
    debt["A16"] = "Total"
    debt["A16"].font = Font(name=FONT, bold=True)
    for cc, L in ((2, "B"), (4, "D"), (8, "H")):
        cell = debt.cell(row=16, column=cc, value=f"=SUM({L}5:{L}14)")
        cell.number_format = BRL
        cell.font = Font(name=FONT, bold=True)
        cell.fill = SOFT
    debt["A18"] = ("'Meses para quitar' supõe juros e parcela constantes: é uma estimativa educativa. "
                   "Confirme as condições com o credor.")
    debt["A18"].font = Font(name=FONT, italic=True, color="6B6B6B")
    for L, wdt in zip("ABCDEFGH", (28, 20, 16, 26, 18, 16, 22, 20)):
        debt.column_dimensions[L].width = wdt
    dv = DataValidation(type="decimal", operator="between", formula1=0, formula2=1, allow_blank=True)
    debt.add_data_validation(dv)
    dv.add("C5:C14")

    sim = wb.create_sheet("Simulador")
    sim["A1"] = "Simulador de quitação"
    sim["A1"].font = Font(name=FONT, bold=True, size=16, color="1F5135")
    rows = [("Saldo da dívida (R$)", 5000, BRL), ("Juros ao mês (%)", 0.05, "0.00%"), ("Quanto vou pagar por mês (R$)", 400, BRL)]
    for i, (lab, val, fmt) in enumerate(rows, 3):
        sim.cell(row=i, column=1, value=lab).font = Font(name=FONT)
        cell = sim.cell(row=i, column=2, value=val)
        cell.fill = YELLOW
        cell.number_format = fmt
        cell.font = Font(name=FONT, color="0000FF")
    sim["A7"] = "Meses para quitar"
    sim["B7"] = '=IFERROR(ROUNDUP(NPER(B4,-B5,B3),0),"Aumente a parcela: ela não cobre os juros")'
    sim["A8"] = "Total pago (R$)"
    sim["B8"] = '=IF(ISNUMBER(B7),B7*B5,"")'
    sim["B8"].number_format = BRL
    sim["A9"] = "Juros pagos no total (R$)"
    sim["B9"] = '=IF(ISNUMBER(B7),B8-B3,"")'
    sim["B9"].number_format = BRL
    sim["A11"] = "Valores de exemplo: troque pelos seus. Estimativa educativa com juros e parcela constantes."
    sim["A11"].font = Font(name=FONT, italic=True, color="6B6B6B")
    for r_ in (7, 8, 9):
        sim.cell(row=r_, column=1).font = Font(name=FONT, bold=True)
        sim.cell(row=r_, column=2).font = Font(name=FONT, bold=True)
        sim.cell(row=r_, column=2).fill = SOFT
    sim.column_dimensions["A"].width = 36
    sim.column_dimensions["B"].width = 24
    wb.save(out)
    return out
